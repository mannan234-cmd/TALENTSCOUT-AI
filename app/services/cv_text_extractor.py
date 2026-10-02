"""
TalentScout AI — CV Text Extraction

Deterministic file-to-text extraction for uploaded CVs.

Supported formats:
- PDF
- DOCX

Public interface:
    cv_text_extractor.extract(
        file_name=file_name,
        file_bytes=file_bytes,
    )
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Optional

from pypdf import PdfReader


class CVTextExtractionError(Exception):
    """Raised when readable CV text cannot be extracted."""


class CVTextExtractor:
    """Extract readable text from supported CV document bytes."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

    def extract(
        self,
        file_name: str,
        file_bytes: bytes,
    ) -> str:
        """Extract readable text from a CV."""

        if not file_name:
            raise CVTextExtractionError(
                "CV filename is required for text extraction."
            )

        if not file_bytes:
            raise CVTextExtractionError(
                "CV file is empty."
            )

        extension = Path(file_name).suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise CVTextExtractionError(
                "Unsupported CV format. Only PDF and DOCX are supported."
            )

        try:
            if extension == ".pdf":
                text = self._extract_pdf(file_bytes)
            else:
                text = self._extract_docx(file_bytes)

        except CVTextExtractionError:
            raise

        except Exception as exc:
            raise CVTextExtractionError(
                f"Unable to extract text from {extension.upper()} CV: {exc}"
            ) from exc

        normalized = self._normalize_text(text)

        if not normalized:
            raise CVTextExtractionError(
                "No readable text was extracted from the CV."
            )

        return normalized

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        """Extract text from PDF bytes using pypdf."""

        try:
            reader = PdfReader(BytesIO(file_bytes))

        except Exception as exc:
            raise CVTextExtractionError(
                f"Unable to read PDF CV: {exc}"
            ) from exc

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):
            try:
                page_text = page.extract_text() or ""

            except Exception as exc:
                raise CVTextExtractionError(
                    f"Unable to extract text from PDF page "
                    f"{page_number}: {exc}"
                ) from exc

            if page_text.strip():
                pages.append(page_text)

        return "\n\n".join(pages)

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        """Extract paragraph and table text from DOCX bytes."""

        try:
            from docx import Document

        except ImportError as exc:
            raise CVTextExtractionError(
                "DOCX extraction requires the python-docx package."
            ) from exc

        try:
            document = Document(BytesIO(file_bytes))

        except Exception as exc:
            raise CVTextExtractionError(
                f"Unable to read DOCX CV: {exc}"
            ) from exc

        parts = []

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()

            if text:
                parts.append(text)

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:
                    cell_text = " ".join(
                        line.strip()
                        for line in (cell.text or "").splitlines()
                        if line.strip()
                    )

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    @staticmethod
    def _normalize_text(
        text: Optional[str]
    ) -> str:
        """Normalize extracted text while preserving line structure."""

        if not text:
            return ""

        normalized_lines = []

        for raw_line in (
            text
            .replace("\r\n", "\n")
            .replace("\r", "\n")
            .split("\n")
        ):
            line = " ".join(raw_line.split())

            if line:
                normalized_lines.append(line)

        return "\n".join(normalized_lines).strip()


cv_text_extractor = CVTextExtractor()
