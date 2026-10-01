from io import BytesIO
from typing import Any, Dict

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class CVExportService:

    @staticmethod
    def _dump(cv: Any) -> Dict[str, Any]:
        if hasattr(cv, "model_dump"):
            return cv.model_dump(mode="json")
        if isinstance(cv, dict):
            return cv
        raise ValueError("Unsupported professional CV object.")

    @classmethod
    def generate_docx(cls, cv: Any) -> bytes:
        data = cls._dump(cv)
        document = Document()

        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(10)

        document.add_heading(
            data.get("full_name") or "Professional CV",
            level=0,
        )

        if data.get("contact_line"):
            p = document.add_paragraph(data["contact_line"])
            p.alignment = 1

        def section(title, value):
            if not value:
                return

            document.add_heading(title, level=1)

            if isinstance(value, str):
                document.add_paragraph(value)
                return

            for item in value:
                if isinstance(item, dict):
                    heading = " — ".join(
                        str(x)
                        for x in (
                            item.get("title", ""),
                            item.get("company", ""),
                        )
                        if x
                    )

                    if item.get("dates"):
                        heading += f" ({item['dates']})"

                    if heading:
                        p = document.add_paragraph()
                        run = p.add_run(heading)
                        run.bold = True

                    if item.get("description"):
                        document.add_paragraph(
                            str(item["description"])
                        )
                else:
                    document.add_paragraph(
                        str(item),
                        style="List Bullet",
                    )

        section(
            "Professional Summary",
            data.get("professional_summary"),
        )
        section("Experience", data.get("experience"))

        if data.get("skills"):
            section(
                "Skills",
                [" • ".join(map(str, data["skills"]))],
            )

        section("Education", data.get("education"))
        section("Certifications", data.get("certifications"))
        section("Projects", data.get("projects"))
        section("Achievements", data.get("achievements"))

        if data.get("languages"):
            section(
                "Languages",
                [" • ".join(map(str, data["languages"]))],
            )

        output = BytesIO()
        document.save(output)
        return output.getvalue()

    @classmethod
    def generate_pdf(cls, cv: Any) -> bytes:
        data = cls._dump(cv)
        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45,
        )

        styles = getSampleStyleSheet()
        story = []

        story.append(
            Paragraph(
                str(data.get("full_name") or "Professional CV"),
                styles["Title"],
            )
        )

        if data.get("contact_line"):
            story.append(
                Paragraph(
                    str(data["contact_line"]),
                    styles["Normal"],
                )
            )

        story.append(Spacer(1, 12))

        def heading(title):
            story.append(
                Paragraph(title, styles["Heading2"])
            )

        def text(value):
            if value:
                story.append(
                    Paragraph(str(value), styles["BodyText"])
                )
                story.append(Spacer(1, 5))

        if data.get("professional_summary"):
            heading("Professional Summary")
            text(data["professional_summary"])

        experience = data.get("experience") or []
        if experience:
            heading("Experience")

            for item in experience:
                title = " — ".join(
                    str(x)
                    for x in (
                        item.get("title", ""),
                        item.get("company", ""),
                    )
                    if x
                )

                if item.get("dates"):
                    title += f" ({item['dates']})"

                if title:
                    story.append(
                        Paragraph(
                            f"<b>{title}</b>",
                            styles["BodyText"],
                        )
                    )

                text(item.get("description"))

        if data.get("skills"):
            heading("Skills")
            text(" • ".join(map(str, data["skills"])))

        for section_name, key in [
            ("Education", "education"),
            ("Certifications", "certifications"),
            ("Projects", "projects"),
            ("Achievements", "achievements"),
            ("Languages", "languages"),
        ]:
            values = data.get(key) or []

            if not values:
                continue

            heading(section_name)

            for item in values:
                if isinstance(item, dict):
                    if item.get("name"):
                        story.append(
                            Paragraph(
                                f"<b>{item['name']}</b>",
                                styles["BodyText"],
                            )
                        )
                    text(item.get("description"))
                else:
                    text(f"• {item}")

        document.build(story)
        return output.getvalue()


cv_export_service = CVExportService()
